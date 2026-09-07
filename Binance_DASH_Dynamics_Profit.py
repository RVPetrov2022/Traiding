# ===================================================================================
# 🤖 АВТОНОМНЫЙ ГИБРИДНЫЙ СЕТОЧНЫЙ РОБОТ С ДИНАМИЧЕСКИМ КАСКАДНЫМ ПРОФИТОМ
# ===================================================================================
#  https://www.perplexity.ai/computer/tasks/b4ef4469-82ab-4534-b7a4-5c07806cd193
#
# После запуска скрипта пытается распродать монеты
#  Если удается распродать все монеты или их часть, то при развороте рынка покупка начинается моне по сетке
# Если рынок сразу пошел вниз, начинается закупка монет по сетке. Не распроданные монеты больше не участвуют в работе скрипта
# После каждой покупки монеты считаеся средняя цена покупки.
# 
#
#
#
#
# ===================================================================================
"""
Что проверено прогоном

Прогнал бота на заглушках Binance в демо-режиме, сценарий «старт 100 → падение до 88 (−12%) → отскок»:

    закуп по 7 уровням сетки, средняя 92.90, брутто-расход 59.95 USDT;

    каскад зафиксировался на 3 частях (глубина 12%), и каждый тейк продал ровно 33.3% базового объёма — раньше первый тейк отдавал бы ~11%;

    после третьего тейка записался CYCLE_COMPLETE, состояние обнулилось, реинвест пересчитал бюджет со 100 до 100.72 USDT;

    рестарт посреди каскада восстановил всё корректно: остаток 0.4302 XAUT, себестоимость 39.97 USDT, брутто-расход круга 59.95, база каскада 0.6453 и сценарий на 3 части.

Что изменилось по конфигу

    GRID_SELL_PROFIT теперь [0.05, 0.10, 0.20, 0.40, 0.60] — как в шапке ТЗ (раньше было до 0.25).

    Добавлена MARKET_REBOUND_PERCENT = 0.05.

    Ключи читаются из BINANCE_API_KEY / BINANCE_API_SECRET; без них реальный режим не стартует.

Перед боевым запуском

Поставьте DEMO_MODE = True и погоняйте пару дней на живом потоке цен — так проверите тайминги сетки и каскада без денег. И обязательно отзовите старые ключи, если они когда-либо попадали в коммиты репозитория.

Две вещи, которые я намеренно не трогал, но их стоит обдумать: комиссия учитывается приблизительно (COMMISSION_RATE вместо фактических fills из ответа биржи), и Excel перечитывается целиком при каждой записи — на длинной истории это начнёт тормозить. Могу переделать оба места, если нужно.
Binance_XAUT_Dynamics_Profit (исправленная версия)
Python
MARKET_REBOUND_PERCENT = 0.05  Это реинвестируется процент от профита ?

Нет, это другое. Легко перепутать — в файле три похожих по духу константы, но они про разное.
Что есть что

MARKET_REBOUND_PERCENT = 0.05 — порог падения рынка для режима «засады» (WAIT_MARKET_DROP). Логика: бот стоит в кэше и ждёт, пока цена упадёт на 5% от якоря, и только тогда включает сетку закупа:

python
target_drop_price = initial_reference_price * (1 - MARKET_REBOUND_PERCENT)
if now_price <= target_drop_price:
    current_direction = "BUYING_GRID"

Реинвестирование профита — это REINVEST_RATE = 0.30, то есть 30% накопленной из Excel прибыли добавляется к базовому депозиту круга:

python
calc_budget = BASE_BUDGET + (total_profit * REINVEST_RATE)
current_cycle_budget = min(round(calc_budget, 2), MAX_TOTAL_BUDGET_USDT)

Именно эту строчку вы видели в прогоне: профит 2.3955 USDT → бюджет вырос со 100 до 100.72.

А GRID_PROFIT = [0.02, ...] — это шаг тейк-профита от средней цены, ещё одна отдельная величина.
Важный нюанс

В вашем коде режим WAIT_MARKET_DROP нигде не включается — ни одна ветка не присваивает current_direction = "WAIT_MARKET_DROP". Стратегия после старта сразу уходит в BUYING_GRID. То есть MARKET_REBOUND_PERCENT сейчас фактически не влияет ни на что: я объявил её только для того, чтобы убрать латентный NameError в мёртвом блоке.

Значение 0.05 я взял как нейтральный дефолт. Если режим засады вам не нужен — можно удалить и блок, и константу. Если нужен (например, чтобы после закрытия цикла бот не выкупал сразу, а ждал просадки) — надо добавить переход в WAIT_MARKET_DROP в месте закрытия цикла. Сказать, как это встроить?
Я так понимаю, запустили скрипт, если цена идет вверх, мы продаем монеты и считаем среднюю цену? Tckb hsyjr gjitk dybp? nj ;ltv jncrjrf d 5 ghjwtynjd jn chtlytq wtys? f gjnjv bltn gjregfnm vjytns gj ctnrt

Порядок примерно такой, но два момента в вашей схеме не совпадают с кодом.
Что происходит при росте

Бот действительно продаёт, но среднюю цену при этом не считает. Первичная сетка продаж (GRID_SELL_RISES = [0.12, 0.15, 0.17, 0.22, 0.25]) отрабатывает от стартового якоря и распродаёт те монеты, что уже лежат у вас на балансе — долями от MAX_XAUT_ACCUMULATION. Бот их не покупал, поэтому себестоимости у них нет, и в Excel в колонку средней пишется просто якорь.

Работает эта сетка только при пустой позиции (total_accumulated_coins < 0.0001) и по сути один раз за историю: флаги САКТИВ больше не сбрасываются. Средняя цена появляется только на покупках.
Что происходит при падении

Здесь главное расхождение: никакого ожидания 5% нет. Бот начинает выкупать сразу, как только цена опустится на 2% от якоря, дальше по всей сетке:

python
GRID_BUY_DROPS = [0.02, 0.03, 0.04, 0.06, 0.08, 0.10, 0.12, 0.15, 0.18]

То есть −2%, −3%, −4%, −6%, −8%, −10%, −12%, −15%, −18% от стартовой цены. После каждой покупки пересчитывается средняя, и когда цена поднимается выше неё — включается каскад тейков.

Про MARKET_REBOUND_PERCENT = 0.05 ещё раз: это не отскок и не от средней. Это порог падения на 5% от якоря в режиме засады, который в вашей текущей логике вообще не активируется.
Полный цикл по факту

Старт → фиксация якоря → одновременно смотрим вверх (разовая распродажа имеющихся монет) и вниз (сетка закупа от −2%) → набрали позицию, посчитали среднюю → цена выше средней → каскад тейков на 1–5 частей по глубине просадки → CYCLE_COMPLETE, новый якорь, бюджет с реинвестом.

Если хотите именно ту логику, которую вы описали — сначала переждать просадку и только потом включать сетку — это как раз тот самый режим засады. Могу его включить: после закрытия цикла бот уходит в WAIT_MARKET_DROP и стартует закуп не сразу, а после падения на заданный процент.
А если мы запустили скрипт и цена пошла вниз, ни одной опреации продажи не произошло, операция покупки монет по сетке будет или бот встанет в ожидании хотя бы одной продажи

Покупки пойдут сразу. Ожидания продаж нет.
Почему

На первом же тике этап старта выставляет режим и сразу отпускает бота в закуп:

python
current_direction = "BUYING_GRID"

Блок первичных продаж (ЭТАП 2) и блок сетки закупа (ЭТАП 4) — это два независимых if. Первый при падении просто не срабатывает (цена ниже якоря, условие now_price >= target_sell_trigger ложно), управление идёт дальше без всяких return. Флаги САКТИВ нигде не участвуют в условиях покупки.

Так что при падении первый ордер уйдёт на уровне −2% от стартовой цены, дальше по сетке −3%, −4% и так далее.
Единственное реальное условие

Наличие свободных USDT на споте. Бюджет берётся из BASE_BUDGET (100 USDT) плюс реинвест, и к продажам он никак не привязан. Если USDT на балансе не хватит, биржа отклонит ордер, флаг шага не выставится, и бот будет пытаться выкупить этот уровень на каждом тике заново.
Побочный эффект, о котором стоит знать

Как только пройдёт первая покупка, total_accumulated_coins станет больше нуля — и первичная сетка продаж отключится до конца цикла, потому что она требует пустой позиции:

python
if executed_sell_steps.count(True) < len(GRID_SELL_VOLUMES) and total_accumulated_coins < 0.0001:

То есть если после старта рынок сразу пошёл вниз, ваши уже лежащие на балансе XAUT по сетке +12%/+15%/+17% в этом круге проданы не будут, даже когда цена потом отрастёт. Дальше распродажей займётся каскад тейков, но он работает только с тем объёмом, который бот купил сам.

"""


import os
import sys
import time
import math
import json
import logging
from datetime import datetime
from openpyxl import load_workbook, Workbook

# Официальные коннекторы Binance
from binance.spot import Spot as Client
from binance.websocket.spot.websocket_stream import SpotWebsocketStreamClient as SpotWebsocketClient



# Инициализация логирования в консоль
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)]
)

# ===================================================================================
# --- РЕЖИМ РАБОТЫ РОБОТА И КЛЮЧИ API ---
# ===================================================================================
DEMO_MODE = False  # True — виртуальный тест в песочнице. False — реальные торги!

# ✅ ПАТЧ №6: Ключи больше НЕ хранятся в исходнике (репозиторий публичный!).
# Задайте их в окружении перед запуском:
#   export BINANCE_API_KEY="..."
#   export BINANCE_API_SECRET="..."
# (или через systemd EnvironmentFile / .env + python-dotenv)

API_KEY = os.getenv("BINANCE_API_KEY", "")
API_SECRET = os.getenv("BINANCE_API_SECRET", "")


if not DEMO_MODE and (not API_KEY or not API_SECRET):
    print("❌ Не заданы BINANCE_API_KEY / BINANCE_API_SECRET в переменных окружения.")
    print("   Запуск в реальном режиме невозможен — все ордера отвалятся по 401.")
    sys.exit(1)

# ===================================================================================
# --- СТРАТЕГИЧЕСКИЕ НАСТРОЕК СЕТОК ДЛЯ МОНЕТЫ XAUT ---
# ===================================================================================
SYMBOL = "XAUTUSDT"         
COMMISSION_RATE = 0.001       # Стандартная комиссия спота Binance (0.1%)

# ➡️ Сетка ПРОДАЖИ (Отрабатывает 1 раз в самый первый запуск, если старт на росте)
GRID_SELL_RISES = [0.12, 0.15, 0.17, 0.22, 0.25]      
GRID_SELL_VOLUMES = [0.1, 0.2, 0.2, 0.2, 0.3]    

# ➡️ Сетка ПОКУПКИ (LONG-Мартингейл сетка донакопления на падающем рынке)
GRID_BUY_DROPS = [0.02, 0.03, 0.04, 0.06, 0.08, 0.10, 0.12, 0.15, 0.18]       
GRID_BUY_VOLUMES = [0.05, 0.05, 0.1, 0.1, 0.1, 0.1, 0.1, 0.20, 0.20]     

# ===================================================================================
# ➡️ ВАША ЛИЧНАЯ ГЕНИАЛЬНАЯ СЕТКА ДИНАМИЧЕСКОГО КАСКАДНОГО ПРОФИТА (ИЗ ТЗ)
# ===================================================================================
GRID_SELL_PROFIT = [0.05, 0.10, 0.20, 0.40, 0.60]      # Пороги падения монеты от фиксации
GRID_PROFIT = [0.02, 0.02, 0.02, 0.02, 0.03]           # Базовый шаг профита для каждого порога
GRID_PART_PROFIT = [1, 2, 3, 4, 5]          # На сколько частей бьем распродажу при отскоке

# ✅ ПАТЧ №1: Константа режима «засады» (WAIT_MARKET_DROP) была использована, но нигде не объявлена.
MARKET_REBOUND_PERCENT = 0.05   # На сколько рынок должен упасть от якоря, чтобы включить сетку закупа 

# ✅ ПАТЧ (описание vs конфиг): пороги ниже соответствуют шапке файла.
# Диапазон падения <5% -> 1 часть, <10% -> 2, <20% -> 3, <40% -> 4, глубже -> 5 частей.


# --- ЖЕСТКИЕ ЛИМИТЫ РИСК-МЕНЕДЖМЕНТА ---
BASE_BUDGET = 100.0             # Базовый стартовый бюджет USDT на цикл подкупа
REINVEST_RATE = 0.30            # Доля реинвестирования чистой прибыли (30% от профита)
MAX_TOTAL_BUDGET_USDT = 300.0   # Жесткий потолок затрат в USDT за один круг
MAX_XAUT_ACCUMULATION = 0.5     # Максимальный лимит удержания монет XAUT на балансе

# Инженерная проверка математики долей при запуске
assert math.isclose(sum(GRID_SELL_VOLUMES), 1.0), "Ошибка: Сумма долей GRID_SELL_VOLUMES должна быть равна 1.0!"
assert math.isclose(sum(GRID_BUY_VOLUMES), 1.0), "Ошибка: Сумма долей GRID_BUY_VOLUMES должна быть равна 1.0!"

# ===================================================================================
# --- СИСТЕМНЫЕ ПУТИ И ДИНАМИЧЕСКОЕ СОСТОЯНИЕ ОЗУ ---
# ===================================================================================
BASE_DIR = os.path.dirname(os.path.abspath(__file__)) if '__file__' in locals() else os.getcwd()
EXCEL_FILE = os.path.join(BASE_DIR, "Binance_XAUT_Dynamic_Profit_mod_v2.xlsx")

current_direction = "WAIT_INITIAL_PRICE"  # Стартовый статус диспетчера стратегии
initial_reference_price = None            # Опорная точка (якорь) для сеток
current_cycle_budget = BASE_BUDGET        # Текущий рабочий бюджет (динамический)

# Сетевые котировки WebSocket
current_websocket_price = None
last_websocket_packet_time = time.time()
last_log_time = 0

# Состояние удерживаемой позиции
total_accumulated_coins = 0.0   # Сколько XAUT сейчас в позиции у робота
current_avg_buy_price = 0.0     # Средняя цена текущей позиции
total_usd_invested = 0.0        # Себестоимость УДЕРЖИВАЕМОГО остатка монет в USDT
max_drop_reached_in_cycle = 0.0 # Пиковый процент падения внутри текущего круга

# ✅ ПАТЧ №5: брутто-расход за круг учитывается ОТДЕЛЬНО от себестоимости остатка.
# Раньше лимит MAX_TOTAL_BUDGET_USDT сбрасывался при каждой продаже и после рестарта.
cycle_usd_spent = 0.0           # Сколько всего USDT потрачено на закупки в текущем круге

# ✅ ПАТЧ №3/№7: снимок объема на входе в фазу тейка + фиксация сценария каскада.
tp_base_volume = 0.0            # Объем позиции на момент первого тейка (база для расчета долей)
tp_locked_idx = None            # Зафиксированный индекс сценария каскада (не меняется посреди распродажи)

# Массивы флагов (Раздельные! Для защиты ОЗУ от ложных совпадений)
executed_buy_steps = [False] * len(GRID_BUY_VOLUMES)
executed_sell_steps = [False] * len(GRID_SELL_VOLUMES)
executed_tp_steps = [False] * 5  # Каскадная маска распродаж (максимум 5 частей из вашего ТЗ)

# --- ИНИЦИАЛИЗАЦИЯ HTTP-КЛИЕНТА Binance ---
binance_client = Client(api_key=API_KEY, api_secret=API_SECRET)
# ===================================================================================
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ ОКРУГЛЕНИЯ И ФИЛЬТРОВ БИРЖИ Binance
# ===================================================================================
def get_symbol_filters():
    """
    Запрашивает актуальные спецификации торговой пары напрямую с биржи Binance.
    Определяет минимальный шаг изменения цены (tickSize) и объема монет (stepSize).
    """
    defaults = {"qty_step": 0.0001, "price_step": 0.01, "min_qty": 0.0001, "min_notional": 5.0}
    if DEMO_MODE:
        return defaults
    try:
        info = binance_client.exchange_info(symbol=SYMBOL)
        # ✅ ПАТЧ №2: "symbols" — ЭТО СПИСОК. Раньше здесь был TypeError,
        # который глушился except — и бот ВСЕГДА работал на дефолтных фильтрах.
        symbols_list = info.get("symbols") or []
        symbol_info = None
        for s in symbols_list:
            if s.get("symbol") == SYMBOL:
                symbol_info = s
                break
        if symbol_info is None:
            print(f"⚠️ Пара {SYMBOL} не найдена в exchange_info. Используем дефолты.")
            return defaults

        qty_step = defaults["qty_step"]
        price_step = defaults["price_step"]
        min_qty = defaults["min_qty"]
        min_notional = defaults["min_notional"]

        for f in symbol_info.get("filters", []):
            ftype = f.get("filterType")
            if ftype == "LOT_SIZE":
                qty_step = float(f["stepSize"])
                min_qty = float(f.get("minQty", min_qty))
            elif ftype == "PRICE_FILTER":
                price_step = float(f["tickSize"])
            # ✅ ПАТЧ: читаем реальный MIN_NOTIONAL биржи вместо хардкода 5.01$
            elif ftype in ("MIN_NOTIONAL", "NOTIONAL"):
                val = f.get("minNotional") or f.get("notional")
                if val is not None:
                    min_notional = float(val)

        print(f"🔧 Фильтры {SYMBOL}: шаг объема {qty_step} | шаг цены {price_step} | мин. объем {min_qty} | мин. сумма {min_notional} USDT")
        return {"qty_step": qty_step, "price_step": price_step, "min_qty": min_qty, "min_notional": min_notional}
    except Exception as e:
        print(f"⚠️ Предупреждение: Не удалось получить фильтры пары с биржи: {e}. Используем дефолты.")
        return defaults

# Инициализируем фильтры (глобальный кэш параметров торговой пары XAUTUSDT)
FILTERS = get_symbol_filters()

def round_step_down(value, step):
    """
    Округление числа строго вниз с шагом, заданным биржей.
    Защищает от превышения баланса при отправке ордеров.
    """
    if step <= 0:
        return value
    return math.floor(round(value / step, 8)) * step


def select_cascade_index(drop_value):
    """
    ✅ ПАТЧ №7: единая точка выбора сценария каскада по глубине просадки.
    Раньше этот блок был скопирован в двух местах с риском расхождения.
    """
    for idx, drop_threshold in enumerate(GRID_SELL_PROFIT):
        if drop_value <= drop_threshold:
            return idx
    return len(GRID_SELL_PROFIT) - 1


def build_tp_configs(cascade_idx, avg_price):
    """
    Формирует список тейк-уровней (доля объема + триггерная цена) для сценария.
    Доли считаются от БАЗОВОГО объема позиции и в сумме дают ровно 1.0.
    """
    parts_count = GRID_PART_PROFIT[cascade_idx]
    base_profit_step = GRID_PROFIT[cascade_idx]

    configs = []
    for step in range(1, parts_count + 1):
        if parts_count in (2, 3):
            # Спецслучай ТЗ: по 33.33%, остаток — на финальный шаг
            percent_qty = 0.3333 if step < parts_count else (1.0 - 0.3333 * (parts_count - 1))
        else:
            percent_qty = 1.0 / parts_count
        configs.append({
            "percent_qty": percent_qty,
            "trigger": avg_price * (1 + (step * base_profit_step))
        })
    return configs, parts_count, base_profit_step


# ===================================================================================
# ИСПОЛНИТОРЫ ТОРГОВЫХ ОРДЕРОВ БИРЖИ Binance (С ПАРСИНГОМ СДЕЛКИ ИЗ СТАКАНА)
# ===================================================================================
def buy_market_order(usd_amount):
    """
    Отправка рыночного ордера на покупку по MARKET на сумму в USDT.
    Возвращает кортеж: (order_id, actual_qty, avg_price) или (None, 0.0, 0.0) в случае ошибки.
    """
    global current_websocket_price
    if DEMO_MODE: 
        mock_price = current_websocket_price if current_websocket_price else 42.0 
        actual_qty = (usd_amount / mock_price) * (1 - COMMISSION_RATE)
        print(f"🔬 [DEMO-BUY] Виртуальная покупка на сумму {usd_amount:.2f} USDT. Начислено: {actual_qty:.4f} XAUT")
        return f"DEMO_BUY_{int(time.time())}", actual_qty, mock_price

    try:
        amount_usd = round(float(usd_amount), 2)
        # Защитный фильтр MIN_NOTIONAL биржи Binance (берется из реальных фильтров пары)
        min_notional = FILTERS.get("min_notional", 5.0) + 0.01
        if amount_usd < min_notional:
            print(f"⚠️ [ПРОПУСК BUY] Сумма {amount_usd}$ ниже минимума биржи ({min_notional:.2f}$).")
            return None, 0.0, 0.0

        usd_str = f"{amount_usd:.2f}"
        
        # ИСПРАВЛЕНО: create_order вместо старого new_order / Использование quoteOrderQty для закупа на USDT
        response = binance_client.create_order(
            symbol=SYMBOL,
            side='BUY',
            type='MARKET',
            quoteOrderQty=usd_str
        )
        
        if response and 'orderId' in response:
            order_id = response.get('orderId')
            executed_qty = float(response.get('executedQty', 0))
            cummulative_quote_qty = float(response.get('cummulativeQuoteQty', 0))
            
            # Рассчитываем реальную чистую среднюю цену исполнения в стакане
            avg_price = cummulative_quote_qty / executed_qty if executed_qty > 0 else 0.0
            
            print(f"✅ Успешный реальный BUY ордер Binance. ID: {order_id} | Потрачено: {cummulative_quote_qty:.2f} USDT | Получено: {executed_qty:.4f} XAUT")
            return order_id, executed_qty, avg_price
            
        print(f"⚠️ Ошибка Binance API при покупке: Неожиданный ответ {response}")
    except Exception as e:
        print(f"⚠️ Отказ маркет BUY-ордера Binance: {e}")
    return None, 0.0, 0.0


def sell_market_order(coin_qty):
    """
    Отправка рыночного ордера на продажу количества монет XAUT.
    Возвращает кортеж: (order_id, actual_spent_qty, avg_price) или (None, 0.0, 0.0)
    """
    global current_websocket_price
    if coin_qty <= 0:
        return None, 0.0, 0.0
        
    if DEMO_MODE: 
        mock_price = current_websocket_price if current_websocket_price else 42.0
        print(f"🔬 [DEMO-SELL] Виртуальная продажа объема {coin_qty:.4f} XAUT по цене ~{mock_price}")
        return f"DEMO_SELL_{int(time.time())}", coin_qty, mock_price

    try:
        # Безопасно округляем количество монет строго под LOT_SIZE биржи для XAUT
        qty_rounded = round_step_down(coin_qty, FILTERS["qty_step"])
        
        decimal_places = max(0, int(round(-math.log10(FILTERS["qty_step"]))))
        qty_str = f"{qty_rounded:.{decimal_places}f}"
        
        if float(qty_str) < FILTERS.get("min_qty", 0.0001):
            print("⚠️ Отмена продажи: объем после округления ниже minQty биржи.")
            return None, 0.0, 0.0

        # ✅ ПАТЧ: проверка MIN_NOTIONAL НА ПРОДАЖЕ — раньше ее не было,
        # и мелкие куски каскада отклонялись биржей с ошибкой.
        if current_websocket_price:
            notional = float(qty_str) * current_websocket_price
            if notional < FILTERS.get("min_notional", 5.0):
                print(f"⚠️ [ПРОПУСК SELL] Объем {qty_str} XAUT ≈ {notional:.2f}$ ниже минимума биржи.")
                return None, 0.0, 0.0

        # ИСПРАВЛЕНО: create_order вместо старого new_order / Использование quantity для продажи монет
        response = binance_client.create_order(
            symbol=SYMBOL,
            side='SELL',
            type='MARKET',
            quantity=qty_str
        )
        
        if response and 'orderId' in response:
            order_id = response.get('orderId')
            executed_qty = float(response.get('executedQty', 0))
            cummulative_quote_qty = float(response.get('cummulativeQuoteQty', 0))
            avg_price = cummulative_quote_qty / executed_qty if executed_qty > 0 else 0.0
            
            print(f"✅ Успешный реальный SELL ордер Binance. ID: {order_id} | Продано: {executed_qty:.4f} XAUT | Получено: {cummulative_quote_qty:.2f} USDT")
            return order_id, executed_qty, avg_price
            
        print(f"⚠️ Ошибка Binance API при продаже: Неожиданный ответ {response}")
    except Exception as e:
        print(f"⚠️ Отказ маркет SELL-ордера Binance: {e}")
    return None, 0.0, 0.0


# ===================================================================================
# ✅ ПАТЧ: ЗАПРОС СВОБОДНОГО БАЛАНСА АКТИВА (для первичной сетки продаж)
# ===================================================================================
def get_free_balance(asset):
    """
    Возвращает свободный баланс актива на споте. Нужен, чтобы не отправлять 
    ордера на продажу монет, которых физически нет на кошельке.
    """
    if DEMO_MODE:
        return MAX_XAUT_ACCUMULATION if asset == "XAUT" else MAX_TOTAL_BUDGET_USDT
    try:
        account_info = binance_client.account()
        for b in account_info.get('balances', []):
            if b.get('asset') == asset:
                return float(b.get('free', 0.0))
    except Exception as e:
        print(f"⚠️ Не удалось запросить баланс {asset}: {e}")
    return 0.0


# ===================================================================================
# РАСЧЕТ РЕИНВЕСТИРОВАНИЯ БЮДЖЕТА (СЛОЖНЫЙ ПРОЦЕНТ ИЗ EXCEL)
# ===================================================================================
def calculate_reinvest_budget():
    """
    Считывает накопленный чистый профит из Excel таблицы и увеличивает базовый депозит.
    Использует оптимизированный read_only режим без лишней нагрузки на жесткий диск.
    """
    global current_cycle_budget
    if not os.path.exists(EXCEL_FILE):
        current_cycle_budget = BASE_BUDGET
        return
    try:
        wb = load_workbook(EXCEL_FILE, read_only=True, data_only=True, keep_links=False)
        if "Транзакции" not in wb.sheetnames:
            current_cycle_budget = BASE_BUDGET
            wb.close()
            return
            
        ws = wb["Транзакции"]
        total_profit = 0.0
        profit_col_idx = None
        
        row_iterator = ws.iter_rows(values_only=True)
        first_row = next(row_iterator, None)
        
        if first_row:
            for idx, cell_value in enumerate(first_row):
                if cell_value == "Профит в USDT":
                    profit_col_idx = idx
                    break
        
        if profit_col_idx is not None:
            for row_values in row_iterator:
                if row_values and len(row_values) > profit_col_idx:
                    val = row_values[profit_col_idx]
                    if val is not None:
                        try:
                            total_profit += float(str(val).replace(',', '.').strip())
                        except (ValueError, TypeError):
                            continue
        wb.close()
        
        if total_profit > 0:
            calc_budget = BASE_BUDGET + (total_profit * REINVEST_RATE)
            current_cycle_budget = min(round(calc_budget, 2), MAX_TOTAL_BUDGET_USDT)
            print(f"📊 [ЛОНГ РЕИНВЕСТ] Общий профицит из Excel: {total_profit:.4f} USDT. Рабочий бюджет круга: {current_cycle_budget} USDT")
        else:
            current_cycle_budget = BASE_BUDGET
            
    except Exception as e:
        print(f"⚠️ Предупреждение: Ошибка автоматического расчета бюджета: {e}")
        current_cycle_budget = BASE_BUDGET
# ===================================================================================
# РАБОТА С ТАБЛИЦАМИ EXCEL (СТРОГОЕ СООТВЕТСТВИЕ КОЛОНОК)
# ===================================================================================
def log_to_excel(date_s, time_s, buy_p=None, buy_q=None, buy_c=None, sell_p=None, sell_q=None, sell_c=None, avg_b=None, profit_u=None, profit_p=None, step_info="", comment=""):
    """Запись транзакций в Excel со строгим выравниванием колонок и защитой структуры ZIP-архива"""
    
    # 13 строго определенных заголовков таблицы (Индексы: 0-12)
    headers = [
        "Дата",                          # 0  (A)
        "Время сделки",                  # 1  (B)
        "Цена покупки (USDT)",           # 2  (C)
        "Количество купленных монет",    # 3  (D)
        "Стоимость покупки (USDT)",      # 4  (E)
        "Средняя цена купленных монет",  # 5  (F)
        "Количество проданных монет",    # 6  (G)
        "Цена продажи",                  # 7  (H)
        "Стоимость продажи (USDT)",      # 8  (I)
        "Профит в USDT",                 # 9  (J)
        "Профит в %",                    # 10 (K)
        "Фиксация шага покупки",         # 11 (L)
        "Комментарий"                    # 12 (M)
    ]
    
    if not os.path.exists(EXCEL_FILE):
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Транзакции"
            ws.append(headers)
            wb.save(EXCEL_FILE)
            print(f"📁 [EXCEL] Создан новый пустой файл истории: {EXCEL_FILE}")
        except Exception as e:
            print(f"❌ Критическая ошибка создания Excel-шаблона: {e}")
            return

    # Вспомогательная мини-функция для безопасного округления чисел
    def safe_round(val, decimals):
        if val is None or val == "":
            return ""
        try:
            return round(float(val), decimals)
        except (ValueError, TypeError):
            return str(val)

    for attempt in range(3):
        try:
            wb = load_workbook(EXCEL_FILE)
            ws = wb["Транзакции"] if "Транзакции" in wb.sheetnames else wb.active
            
            # СТРОЖАЙШЕЕ выравнивание данных по индексам заголовков таблицы
            row_data = [
                date_s,                               # A: Дата
                time_s,                               # B: Время сделки
                safe_round(buy_p, 2),                 # C: Цена покупки (USDT)
                safe_round(buy_q, 4),                 # D: Количество купленных монет
                safe_round(buy_c, 2),                 # E: Стоимость покупки (USDT)
                safe_round(avg_b, 2),                 # F: Средняя цена купленных монет
                safe_round(sell_q, 4),                # G: Количество проданных монет
                safe_round(sell_p, 2),                # H: Цена продажи
                safe_round(sell_c, 2),                # I: Стоимость продажи (USDT)
                safe_round(profit_u, 4),              # J: Профит в USDT
                safe_round(profit_p, 2),              # K: Профит в %
                str(step_info),                       # L: Фиксация шага покупки
                str(comment)                          # M: Комментарий
            ]
            
            ws.append(row_data)
            wb.save(EXCEL_FILE)
            
            # --- БЕЗОПАСНАЯ ХАРДВЕРНАЯ ЗАЩИТА LINUX (ДЛЯ USB-ФЛЕШЕК / КЭША VPS) ---
            if sys.platform.startswith('linux'):
                os.sync()  # Атомарный системный flush для сохранения стабильности файловой системы
            break
        except PermissionError:
            print(f"⚠️ [БЛОКИРОВКА EXCEL] Ожидание 3 сек (Попытка {attempt+1}/3)... Таблица открыта в LibreOffice.")
            time.sleep(3)
        except Exception as excel_err:
            print(f"❌ Сбой записи Excel: {excel_err}")
            break


# ===================================================================================
# ВОССТАНОВЛЕНИЕ СОСТОЯНИЯ (НАЧАЛО МОДУЛЯ И СИНХРОНИЗАЦИЯ С ДИСКОМ)
# ===================================================================================
def restore_state_from_excel():
    """
    Восстанавливает состояние шагов сетки, каскадных тейков и среднюю цену из Excel.
    Полностью исключает баг пустых строк max_row и ложные подкупы при рестарте в screen.
    """
    global current_direction, initial_reference_price, total_accumulated_coins, current_avg_buy_price
    global executed_buy_steps, executed_sell_steps, executed_tp_steps, max_drop_reached_in_cycle
    global total_usd_invested, cycle_usd_spent, tp_base_volume, tp_locked_idx
    
    if not os.path.exists(EXCEL_FILE):
        print("📝 Истории торгов на диске не обнаружено. Включаем чистый АВТОНОМНЫЙ режим.")
        return False
    try:
        # Читаем только значения, закрываем ссылки для экономии ОЗУ
        wb = load_workbook(EXCEL_FILE, data_only=True, keep_links=False)
        if "Транзакции" not in wb.sheetnames:
            wb.close()
            return False
        ws = wb["Транзакции"]
        
        # Загружаем все строки в память (избегаем бага пустых строк)
        all_rows = list(ws.iter_rows(values_only=True))
        wb.close()  
        
        if len(all_rows) < 2:  
            return False
            
        header = list(all_rows[0])
        try:
            idx_comment = header.index("Комментарий")
            idx_step_inf = header.index("Фиксация шага покупки")
            idx_avg_p = header.index("Средняя цена купленных монет")
            idx_buy_p = header.index("Цена покупки (USDT)")
            idx_buy_q = header.index("Количество купленных монет")
            idx_sell_q = header.index("Количество проданных монет")
            idx_buy_c = header.index("Стоимость покупки (USDT)")  # ✅ ПАТЧ №5
        except ValueError as e:
            print(f"⚠️ Ошибка: В структуре Excel файла не найдены нужные колонки: {e}")
            return False

        # Фильтруем пустые строки с конца таблицы
        real_rows = [r for r in all_rows[1:] if any(c is not None and str(c).strip() != "" for c in r)]
        if not real_rows:
            return False
            
        # Дальнейший парсинг и сборка масок переходит в Часть 4...
        last_row = real_rows[-1]
        comment = str(last_row[idx_comment] or "")
        step_inf = str(last_row[idx_step_inf] or "")
        
        # --- ЖЕЛЕЗОБЕТОННОЕ ВОССТАНОВЛЕНИЕ ФЛАГОВ СЕТКИ И КАСКАДА ТЕЙКОВ ИЗ МАСКИ ---
        if "БАКТИВ:" in comment and "САКТИВ:" in comment:
            try:
                parts = comment.split("|")
                for p in parts:
                    if "БАКТИВ:" in p:
                        b_str = p.split("БАКТИВ:")[-1].strip()
                        parsed_buy = [char == '1' for char in b_str]
                        for i in range(min(len(executed_buy_steps), len(parsed_buy))):
                            executed_buy_steps[i] = parsed_buy[i]
                            
                    if "САКТИВ:" in p:
                        s_str = p.split("САКТИВ:")[-1].strip()
                        parsed_sell = [char == '1' for char in s_str]
                        for i in range(min(len(executed_sell_steps), len(parsed_sell))):
                            executed_sell_steps[i] = parsed_sell[i]
                            
                    # Чтение флагов каскадных Тейк-Профитов из Excel (Имена строго по Части 1)
                    if "ТПАКТИВ:" in p:
                        tp_str = p.split("ТПАКТИВ:")[-1].strip()
                        parsed_tp = [char == '1' for char in tp_str]
                        for i in range(min(len(executed_tp_steps), len(parsed_tp))):
                            executed_tp_steps[i] = parsed_tp[i]
            except Exception as e:
                print(f"⚠️ Предупреждение: Ошибка восстановления маски шагов из комментария: {e}")

        # Функция безопасной конвертации данных из Excel-ячеек в float
        def safe_float(val):
            if val is None or str(val).strip() == "":
                return 0.0
            try:
                return float(str(val).replace(',', '.').strip())
            except ValueError:
                return 0.0

        # Ищем последнюю валидную среднюю цену закупки (идем снизу вверх)
        current_avg_buy_price = 0.0
        for r in reversed(real_rows):
            val = r[idx_avg_p]
            if val is not None:
                price_converted = safe_float(val)
                if price_converted > 0:
                    current_avg_buy_price = price_converted
                    break
                
        # --- ИСПРАВЛЕННЫЙ РАСЧЕТ ТЕКУЩЕГО НАКОПЛЕННОГО ПУЛА МОНЕТ ДЛЯ ТЕКУЩЕГО ЦИКЛА ---
        # Ищем точку последнего полного закрытия сетки (CYCLE_COMPLETE) или сброса (RESET)
        start_pool_idx = 0
        for idx, r in enumerate(reversed(real_rows)):
            row_step = str(r[idx_step_inf] or "")
            if "CYCLE_COMPLETE" in row_step or "RESET" in row_step:
                # Нашли предыдущий закрытый круг. Считаем баланс строго от этой точки вниз
                start_pool_idx = len(real_rows) - idx
                break

        coins_pool = 0.0
        spent_pool = 0.0
        for r in real_rows[start_pool_idx:]:
            coins_pool += safe_float(r[idx_buy_q])
            coins_pool -= safe_float(r[idx_sell_q])
            spent_pool += safe_float(r[idx_buy_c])   # ✅ ПАТЧ №5: брутто-расход круга
                
        # Округляем до 4 знаков в соответствии с точностью XAUT
        total_accumulated_coins = round(max(0.0, coins_pool), 4)
        cycle_usd_spent = round(max(0.0, spent_pool), 2)

        # ✅ ПАТЧ №5 (главное): раньше total_usd_invested НЕ восстанавливался и оставался 0.
        # Первый же подкуп после рестарта давал заниженную среднюю цену и тейки ниже безубытка.
        # Себестоимость остатка = средняя цена × оставшиеся монеты.
        if current_avg_buy_price > 0 and total_accumulated_coins > 0:
            total_usd_invested = round(current_avg_buy_price * total_accumulated_coins, 2)
        else:
            total_usd_invested = 0.0

        # --- ВОССТАНОВЛЕНИЕ СТАРТОВОГО ЯКОРЯ ---
        # Ищем строку START или RESET текущего неоконченного цикла для определения первоначальной цены
        initial_reference_price = None
        for r in reversed(real_rows[start_pool_idx:]):
            row_step = str(r[idx_step_inf] or "")
            if "START" in row_step or "RESET" in row_step:
                initial_reference_price = safe_float(r[idx_avg_p])
                break

        if initial_reference_price is None or initial_reference_price == 0:
            initial_reference_price = current_avg_buy_price if current_avg_buy_price > 0 else safe_float(last_row[idx_avg_p])

        # --- РАСЧЕТ МАКСИМАЛЬНОГО ДОСТИГНУТОГО ПАДЕНИЯ В ТЕКУЩЕМ КРУГЕ ---
        if initial_reference_price > 0:
            min_buy_p = initial_reference_price
            for r in real_rows[start_pool_idx:]:
                b_p = safe_float(r[idx_buy_p])
                if 0 < b_p < min_buy_p:
                    min_buy_p = b_p
            # Переменная зафиксирует максимальный шаг падения, который бот пролетел до рестарта
            max_drop_reached_in_cycle = (initial_reference_price - min_buy_p) / initial_reference_price

        # --- ✅ ПАТЧ №3/№7: ВОССТАНОВЛЕНИЕ БАЗЫ КАСКАДА ПОСЛЕ РЕСТАРТА ---
        # Если часть тейков уже исполнена, восстанавливаем исходный объем фазы тейка
        # (остаток + все проданное в рамках каскада) и фиксируем сценарий каскада.
        tp_base_volume = 0.0
        tp_locked_idx = None
        if executed_tp_steps.count(True) > 0:
            sold_in_cascade = 0.0
            for r in real_rows[start_pool_idx:]:
                if "TP_CASC_" in str(r[idx_step_inf] or ""):
                    sold_in_cascade += safe_float(r[idx_sell_q])
            tp_base_volume = round(total_accumulated_coins + sold_in_cascade, 4)
            tp_locked_idx = select_cascade_index(max_drop_reached_in_cycle)

        # --- МАТРИЦА НАПРАВЛЕНИЯ ДВИЖЕНИЯ РОБОТА ДЛЯ ИСКЛЮЧЕНИЯ ЛОЖНЫХ ВХОДОВ ---
        if "START" in step_inf:
            current_direction = "BUYING_GRID"  # Сразу уходим на одновременный сквозной мониторинг сеток
        elif "BUY_" in step_inf or "ACTIVATE_BUY" in step_inf:
            current_direction = "BUYING_GRID"
        elif "WAIT_TAKE" in step_inf or "TP_CASC_" in step_inf or total_accumulated_coins > 0:
            current_direction = "WAIT_TAKE_PROFIT"
        else:
            current_direction = "BUYING_GRID"
        print(f"🛡️ [СИНХРОНИЗАЦИЯ] Массив шагов успешно восстановлен.")
        print(f"    Исполнено BUY: {executed_buy_steps.count(True)} | Исполнено SELL: {executed_sell_steps.count(True)} | Частичных ТП: {executed_tp_steps.count(True)}")
        print(f"    Позиция в ОЗУ: {total_accumulated_coins:.4f} XAUT по средней цене {current_avg_buy_price:.2f} USDT. РЕЖИМ: {current_direction}")
        print(f"    Себестоимость остатка: {total_usd_invested:.2f} USDT | Потрачено за круг: {cycle_usd_spent:.2f} USDT")
        print(f"    Зафиксировано макс. падение в текущем цикле: {max_drop_reached_in_cycle * 100:.2f}%")
        return True
    except Exception as e:
        print(f"⚠️ Сбой синхронизации Excel при восстановлении состояния: {e}")
    return False


# ===================================================================================
# ЗАПРОС РЕАЛЬНЫХ БАЛАНСОВ С БИРЖИ Binance
# ===================================================================================
def show_real_balances():
    """
    Запрашивает реальные балансы спотового кошелька на Binance с защитой от ошибок.
    Считывает как свободные средства, так и замороженные в лимитных ордерах.
    """
    if DEMO_MODE: 
        return
    try:
        # ИСПРАВЛЕНО: В официальном binance.spot используется метод account(), а не get_account()
        account_info = binance_client.account()
        
        # Извлекаем список балансов
        balances = account_info.get('balances', [])
        
        usdt_free = 0.0
        usdt_locked = 0.0
        xaut_free = 0.0
        xaut_locked = 0.0
        
        # Парсим данные по целевым активам
        for b in balances:
            asset = b.get('asset')
            if asset == 'USDT':
                usdt_free = float(b.get('free', 0.0))
                usdt_locked = float(b.get('locked', 0.0))
            elif asset == 'XAUT':
                xaut_free = float(b.get('free', 0.0))
                xaut_locked = float(b.get('locked', 0.0))
        
        total_usdt = usdt_free + usdt_locked
        total_xaut = xaut_free + xaut_locked
        
        print("-" * 75)
        print(f"💰 [БАНКОВСКИЙ ОТЧЕТ БИРЖИ] Кошелек синхронизирован:")
        print(f"    USDT -> Свободно: {usdt_free:>10.2f} $ | В ордерах: {usdt_locked:.2f} $ | Всего: {total_usdt:.2f} $")
        print(f"    XAUT -> Свободно: {xaut_free:>10.4f}   | В ордерах: {xaut_locked:.4f}   | Всего: {total_xaut:.4f}")
        print("-" * 75)
        
        # Защитный кросс-чек на случай ручного вмешательства трейдера
        if xaut_locked > 0 and total_accumulated_coins == 0:
            print("⚠️ ВНИМАНИЕ: На бирже обнаружены замороженные монеты XAUT, но локальный цикл робота пуст!")
        
    except Exception as e:
        print(f"⚠️ Не удалось обновить балансы с биржи Binance: {e}")

# ===================================================================================
# ПРИЕМНИК ДАННЫХ WEBSOCKET (ИСПРАВЛЕН ПОД СИГНАТУРУ ВАШЕЙ БИБЛИОТЕКИ)
# ===================================================================================
def handle_ticker_message(ws, message):
    """
    Приемник данных: принимает два аргумента (ws, message) в соответствии с Вашим SDK.
    Идеально парсит цены из официального потока @miniTicker.
    """
    global current_websocket_price, last_websocket_packet_time
    try:
        # Библиотека может передавать как строку, так и готовый словарь
        if isinstance(message, dict): 
            msg = message
        else: 
            msg = json.loads(message)
        
        # --- ПАРСИНГ ДАННЫХ ДЛЯ КАСКАДНОЙ СТРАТЕГИИ ---
        # 1. Если данные пришли в формате мини-тикера (ключ "e" равен "24hrMiniTicker")
        if isinstance(msg, dict) and msg.get("e") == "24hrMiniTicker" and msg.get("s") == SYMBOL:
            if "c" in msg:  # В мини-тикере текущая цена лежит строго в ключе "c"
                current_websocket_price = float(msg["c"])
                last_websocket_packet_time = time.time()
                return

        # 2. Универсальный резервный парсинг для обычного тикера
        elif isinstance(msg, dict) and msg.get("s") == SYMBOL:
            if "c" in msg:
                current_websocket_price = float(msg["c"])
                last_websocket_packet_time = time.time()
                return
            elif "lastPrice" in msg:
                current_websocket_price = float(msg["lastPrice"])
                last_websocket_packet_time = time.time()
                return

    except Exception as e:
        print(f"❌ Ошибка внутри функции handle_ticker_message: {e}")


# ===================================================================================
# УПРАВЛЯЮЩИЙ ДВИЖОК ТОРГОВОЙ СТРАТЕГИИ (ЧАСТЬ А: ИНИЦИАЛИЗАЦИЯ КРУГА)
# ===================================================================================
def run_trading_strategy_step(now_price):
    global current_direction, initial_reference_price, total_accumulated_coins, current_avg_buy_price
    global total_usd_invested, executed_buy_steps, executed_sell_steps, executed_tp_steps, last_log_time 
    global current_cycle_budget, max_drop_reached_in_cycle
    global cycle_usd_spent, tp_base_volume, tp_locked_idx
    
    dt = datetime.now()
    date_str, time_str = dt.strftime("%Y-%m-%d"), dt.strftime("%H:%M:%S")

    # ===================================================================================
    # ЭТАП 1: ЧИСТЫЙ СТАРТ СИСТЕМЫ (Выполняется 1 раз, если Excel пустой или при сбросе)
    # ===================================================================================
    if current_direction == "WAIT_INITIAL_PRICE" or initial_reference_price is None:
        initial_reference_price = now_price
        current_avg_buy_price = now_price
        current_direction = "BUYING_GRID"  # Скрипт сразу уходит на параллельный мониторинг сеток покупок/продаж
        total_accumulated_coins = 0.0      # Начинаем с чистого кэша USDT
        total_usd_invested = 0.0
        cycle_usd_spent = 0.0
        max_drop_reached_in_cycle = 0.0
        tp_base_volume = 0.0
        tp_locked_idx = None
        executed_buy_steps = [False] * len(GRID_BUY_VOLUMES)
        executed_sell_steps = [False] * len(GRID_SELL_VOLUMES)
        executed_tp_steps = [False] * 5
        
        log_to_excel(
            date_s=date_str, time_s=time_str, avg_b=now_price, step_info="START", 
            comment="ФИКСАЦИЯ СТАРТОВОЙ ЦЕНЫ. ОДНОВРЕМЕННЫЙ МОНИТОРИНГ СЕТОК."
        )
        print(f"🎯 Точка старта успешно зафиксирована: {initial_reference_price:.2f} USDT. Робот включил сканирование биржи...")
        return

    # ✅ ПАТЧ №8: глубина просадки теперь трекается НА КАЖДОМ ТИКЕ, а не только в момент покупки.
    # Раньше при исчерпанном бюджете дальнейшее падение не учитывалось и каскад выбирал
    # слишком мягкий сценарий. Во время активного каскада сценарий уже заблокирован (tp_locked_idx).
    if initial_reference_price and initial_reference_price > 0:
        tick_drop = (initial_reference_price - now_price) / initial_reference_price
        if tick_drop > max_drop_reached_in_cycle:
            max_drop_reached_in_cycle = tick_drop

    b_state = "".join(["1" if x else "0" for x in executed_buy_steps])
    s_state = "".join(["1" if x else "0" for x in executed_sell_steps])
    tp_state = "".join(["1" if x else "0" for x in executed_tp_steps])
    status_comment = f"БАКТИВ:{b_state}|САКТИВ:{s_state}|ТПАКТИВ:{tp_state}"

    # ===================================================================================
    # ЭТАП 2: СИНХРОННАЯ СЕТКА ПЕРВИЧНЫХ ПРОДАЖ (ОТРАБАТЫВАЕТ ТОЛЬКО 1 РАЗ ЗА ВСЮ ИСТОРИЮ)
    # ===================================================================================
    # Если на руках еще нет купленных монет — проверяем уровни роста для сбора USDT
    if executed_sell_steps.count(True) < len(GRID_SELL_VOLUMES) and total_accumulated_coins < 0.0001:
        for i in range(len(GRID_SELL_RISES)):
            target_sell_trigger = initial_reference_price * (1 + GRID_SELL_RISES[i])
            if now_price >= target_sell_trigger and not executed_sell_steps[i]:
                
                # Доля монет на продажу (требует наличия XAUT на балансе)
                coins_portion = MAX_XAUT_ACCUMULATION * GRID_SELL_VOLUMES[i]

                # ✅ ПАТЧ: раньше бот пытался продать монеты, которых могло не быть на кошельке,
                # и получал серию отказов биржи. Теперь сверяемся с реальным свободным балансом.
                free_xaut = get_free_balance("XAUT")
                if free_xaut < coins_portion:
                    if free_xaut < FILTERS.get("min_qty", 0.0001):
                        print(f"ℹ️ [Первичная Продажа] Шаг {i+1} пропущен: свободных XAUT на балансе нет.")
                        executed_sell_steps[i] = True  # Не долбимся в биржу каждую секунду
                        break
                    print(f"⚠️ [Первичная Продажа] Урезание объема шага {i+1}: {coins_portion:.4f} -> {free_xaut:.4f} XAUT")
                    coins_portion = free_xaut

                order_id, actual_sell_qty, actual_sell_price = sell_market_order(coins_portion)
                
                if order_id and actual_sell_qty > 0:
                    executed_sell_steps[i] = True
                    usd_received = actual_sell_qty * actual_sell_price
                    
                    log_to_excel(
                        date_s=date_str, time_s=time_str, avg_b=initial_reference_price,
                        sell_q=actual_sell_qty, sell_p=actual_sell_price, sell_c=usd_received,
                        step_info=f"FIRST_S_{i+1}", comment=f"ПЕРВИЧНАЯ РАСПРОДАЖА ШАГ {i+1} | {status_comment}"
                    )
                    print(f"💰 [Первичная Продажа] Исполнен шаг {i+1}. Цена: {actual_sell_price:.2f} USDT")
                    break

    # Логика стратегии переходит в Часть 7 (Подкуп по сетке и динамический каскад)...
    # ===================================================================================
    # ЭТАП 3: РЕЖИМ ЗАСАДЫ (Ждем снижения рынка после полной фиксации прибыли)
    # ===================================================================================
    if current_direction == "WAIT_MARKET_DROP":
        target_drop_price = initial_reference_price * (1 - MARKET_REBOUND_PERCENT)
        if now_price <= target_drop_price:
            current_direction = "BUYING_GRID"
            initial_reference_price = now_price  # Текущая цена — новая опора (якорь) засады
            executed_buy_steps = [False] * len(GRID_BUY_VOLUMES)
            total_usd_invested = 0.0
            cycle_usd_spent = 0.0
            total_accumulated_coins = 0.0
            max_drop_reached_in_cycle = 0.0
            tp_base_volume = 0.0
            tp_locked_idx = None
            print(f"📉 Рынок упал до {now_price:.2f} USDT. АКТИВИРУЕМ СЕТКУ ПОКУПКИ.")
            log_to_excel(
                date_s=date_str, time_s=time_str, buy_p=now_price, avg_b=initial_reference_price, step_info="ACTIVATE_BUY",
                comment="РЫНОК ДОСТИГ ТОЧКИ СНИЖЕНИЯ. СЕТКА ПОДКУПА ВКЛЮЧЕНА."
            )

    # ===================================================================================
    # ЭТАП 4: НАКОПЛЕНИЕ ПО СЕТКЕ ПОКУПКИ (С КОРРЕКЦИЕЙ НА СБОИ СЕТИ)
    # ===================================================================================
    if current_direction == "BUYING_GRID":
        # Сканируем всю сетку. Если из-за сбоя проскочили НЕСКОЛЬКО уровней,
        # этот цикл последовательно за несколько тиков выкупит их все, не пропуская объемы.
        for i in range(len(GRID_BUY_DROPS)):
            target_buy_trigger = initial_reference_price * (1 - GRID_BUY_DROPS[i])
            
            if now_price <= target_buy_trigger and not executed_buy_steps[i]:
                # ✅ ПАТЧ №5: лимит считается по БРУТТО-расходу круга (cycle_usd_spent),
                # а не по себестоимости остатка — иначе лимит сбрасывался после каждой продажи.
                if cycle_usd_spent >= MAX_TOTAL_BUDGET_USDT:
                    print(f"⚠️ [РИСК-МЕНЕДЖМЕНТ] Бюджет круга исчерпан ({cycle_usd_spent:.2f} / {MAX_TOTAL_BUDGET_USDT} USDT).")
                    break

                # ✅ ПАТЧ: лимит MAX_XAUT_ACCUMULATION теперь реально проверяется (раньше был декларативным).
                if total_accumulated_coins >= MAX_XAUT_ACCUMULATION:
                    print(f"⚠️ [РИСК-МЕНЕДЖМЕНТ] Достигнут потолок удержания {MAX_XAUT_ACCUMULATION} XAUT.")
                    break

                usd_allocation = current_cycle_budget * GRID_BUY_VOLUMES[i]
                if cycle_usd_spent + usd_allocation > MAX_TOTAL_BUDGET_USDT:
                    usd_allocation = MAX_TOTAL_BUDGET_USDT - cycle_usd_spent

                # Не даем пробить потолок монет одним ордером
                room_coins = MAX_XAUT_ACCUMULATION - total_accumulated_coins
                if now_price > 0 and usd_allocation / now_price > room_coins:
                    usd_allocation = room_coins * now_price

                order_id, actual_qty, actual_price = buy_market_order(usd_allocation)
                if order_id and actual_qty > 0:
                    executed_buy_steps[i] = True
                    
                    # Фиксируем максимальный процент падения от стартовой цены для динамического каскада
                    current_drop = (initial_reference_price - actual_price) / initial_reference_price
                    if current_drop > max_drop_reached_in_cycle:
                        max_drop_reached_in_cycle = current_drop
                    
                    fact_usd_spent = actual_qty * actual_price
                    coins_net = actual_qty * (1 - COMMISSION_RATE)
                    
                    total_usd_invested = round(total_usd_invested + fact_usd_spent, 2)
                    cycle_usd_spent = round(cycle_usd_spent + fact_usd_spent, 2)
                    total_accumulated_coins = round(total_accumulated_coins + coins_net, 4)
                    current_avg_buy_price = round(total_usd_invested / total_accumulated_coins, 2)
                    
                    log_to_excel(
                        date_s=date_str, time_s=time_str, buy_p=actual_price, buy_q=coins_net, buy_c=fact_usd_spent,
                        avg_b=current_avg_buy_price, step_info=f"BUY_{i+1}", comment=f"ПОДКУП НА ШАГЕ {i+1} (МаксСниж: {max_drop_reached_in_cycle*100:.1f}%) | {status_comment}"
                    )
                    print(f"🛒 [ПОКУПКА] Шаг {i+1} исполнен. Ср.цена закупки: {current_avg_buy_price:.2f} USDT")
                    show_real_balances()
                    return # Даем системе обновить данные на следующем тике цены

        # Выход в режим ожидания продаж, если монеты куплены и цена пошла вверх
        if total_accumulated_coins > 0.0001 and now_price > current_avg_buy_price:
            current_direction = "WAIT_TAKE_PROFIT"
            print(f"🚀 Цена превысила среднюю закупку ({current_avg_buy_price:.2f}). Включаем каскад Тейка.")

    # ===================================================================================
    # ЭТАП 5: РЕЖИМ ГЕНИАЛЬНОГО КАСКАДНОГО ТЕЙК-ПРОФИТА И ЖЕСТКОГО СБРОСА ЦИКЛА
    # ===================================================================================
    elif current_direction == "WAIT_TAKE_PROFIT":
        
        # ⚡ КЛЮЧЕВОЙ ДИНАМИЧЕСКИЙ РЕВЕРС / СБРОС ЗАСТРЯВШИХ МОНЕТ ДРУГОМУ БОТУ (Пункт 5 ТЗ)
        if now_price < current_avg_buy_price:
            if executed_tp_steps.count(True) > 0:
                # Согласно Вашему ТЗ: мы совершили часть продаж, но рынок снова пошел вниз. 
                # "Забываем" про остаток монет, жестко сбрасываем цикл и начинаем круг с нуля!
                log_to_excel(
                    date_s=date_str, time_s=time_str, avg_b=current_avg_buy_price, step_info="RESET_CYCLE_ABANDON",
                    comment=f"⚠️ РЕВЕРС ВНИЗ ПОСЛЕ ЧАСТИЧНОГО ТЕЙКА. Остаток {total_accumulated_coins:.4f} XAUT оставлен. НАЧАЛО С ЧИСТОГО ЛИСТА."
                )
                print("⚠️ [СБРОС ЦИКЛА] Цена упала ниже средней после частичного Тейка. Оставляем остаток другому боту и начинаем круг заново!")
                
                # Полная очистка ОЗУ под новый круг закупа
                initial_reference_price = now_price
                current_avg_buy_price = now_price
                current_direction = "BUYING_GRID"
                total_accumulated_coins = 0.0
                total_usd_invested = 0.0
                cycle_usd_spent = 0.0
                max_drop_reached_in_cycle = 0.0
                tp_base_volume = 0.0
                tp_locked_idx = None
                executed_buy_steps = [False] * len(GRID_BUY_VOLUMES)
                executed_tp_steps = [False] * 5  
                calculate_reinvest_budget()
                return
            else:
                # Если частичных продаж еще не было, это стандартный донабор позиции по нижней сетке
                current_direction = "BUYING_GRID"
                print("⚠️ Цена ниже средней закупки. Возврат в режим донакопления позиций.")
                return

        # --- ✅ ПАТЧ №7: СЦЕНАРИЙ КАСКАДА ФИКСИРУЕТСЯ ОДИН РАЗ НА ФАЗУ РАСПРОДАЖИ ---
        # Раньше при углублении просадки сетка пересобиралась (было 2 части — стало 4),
        # а маска executed_tp_steps оставалась от старого сценария — доли ехали.
        if tp_locked_idx is None:
            tp_locked_idx = select_cascade_index(max_drop_reached_in_cycle)
            print(f"🔒 Сценарий каскада зафиксирован: {GRID_PART_PROFIT[tp_locked_idx]} частей (глубина {max_drop_reached_in_cycle*100:.1f}%)")

        # --- ✅ ПАТЧ №3: БАЗОВЫЙ ОБЪЕМ ФАЗЫ ТЕЙКА ФИКСИРУЕТСЯ ОДИН РАЗ ---
        if tp_base_volume <= 0.0:
            tp_base_volume = total_accumulated_coins

        tp_configs, parts_count, base_profit_step = build_tp_configs(tp_locked_idx, current_avg_buy_price)

        # --- ИСПОЛНЕНИЕ КАСКАДНЫХ ТЕЙК-ПРОФИТОВ ---
        for idx, tp in enumerate(tp_configs):
            if now_price >= tp["trigger"] and not executed_tp_steps[idx]:
                # ✅ ПАТЧ №3: доля считается от ЗАФИКСИРОВАННОЙ базы фазы тейка.
                # Раньше объем резался ДВАЖДЫ (остаток / число шагов × доля):
                # при 3 частях первый тейк продавал ~11% позиции вместо 33%.
                if idx == len(tp_configs) - 1:
                    sell_qty = total_accumulated_coins  # Финальный шаг забирает весь остаток
                else:
                    sell_qty = round(min(tp_base_volume * tp["percent_qty"], total_accumulated_coins), 4)
                
                if sell_qty <= 0.0:
                    continue
                    
                order_id, actual_sell_qty, actual_sell_price = sell_market_order(sell_qty)
                if order_id and actual_sell_qty > 0:
                    executed_tp_steps[idx] = True
                    
                    usd_received = actual_sell_qty * actual_sell_price
                    clean_usd_received = usd_received * (1 - COMMISSION_RATE)
                    
                    # Пропорционально уменьшаем учтенную себестоимость остатка
                    allocated_usd_share = (actual_sell_qty / total_accumulated_coins) * total_usd_invested if total_accumulated_coins > 0 else total_usd_invested
                    net_profit = clean_usd_received - allocated_usd_share
                    
                    total_accumulated_coins = round(max(0.0, total_accumulated_coins - actual_sell_qty), 4)
                    total_usd_invested = round(max(0.0, total_usd_invested - allocated_usd_share), 2)

                    # Обновляем маску СРАЗУ после сделки, чтобы в Excel попало актуальное состояние
                    fresh_status = (
                        f"БАКТИВ:{''.join('1' if x else '0' for x in executed_buy_steps)}"
                        f"|САКТИВ:{''.join('1' if x else '0' for x in executed_sell_steps)}"
                        f"|ТПАКТИВ:{''.join('1' if x else '0' for x in executed_tp_steps)}"
                    )
                    
                    log_to_excel(
                        date_s=date_str, time_s=time_str, avg_b=current_avg_buy_price,
                        sell_q=actual_sell_qty, sell_p=actual_sell_price, sell_c=clean_usd_received,
                        profit_u=net_profit, profit_p=(net_profit / (allocated_usd_share if allocated_usd_share > 0 else 1)) * 100,
                        step_info=f"TP_CASC_{idx+1}", comment=f"💰 КАСКАДНЫЙ ТЕЙК №{idx+1}/{parts_count} (Глубина: {max_drop_reached_in_cycle*100:.1f}%) | {fresh_status}"
                    )
                    print(f"💰 [ТЕЙК №{idx+1}/{parts_count}] Продано {actual_sell_qty:.4f} XAUT по {actual_sell_price:.2f} | Чистый профит: {net_profit:+.4f} USDT")
                    show_real_balances()

                    # --- ✅ ПАТЧ №4: КОРРЕКТНОЕ ЗАКРЫТИЕ ЦИКЛА ---
                    # Раньше после последнего тейка бот оставался в WAIT_TAKE_PROFIT с нулем монет,
                    # маркер CYCLE_COMPLETE не писался (а восстановление его ищет!),
                    # и на первом же тике ниже средней цикл уходил в RESET_CYCLE_ABANDON.
                    all_tp_done = all(executed_tp_steps[k] for k in range(len(tp_configs)))
                    if all_tp_done or total_accumulated_coins < FILTERS.get("min_qty", 0.0001):
                        log_to_excel(
                            date_s=date_str, time_s=time_str, avg_b=now_price, step_info="CYCLE_COMPLETE",
                            comment=f"✅ КРУГ ПОЛНОСТЬЮ ЗАКРЫТ. Остаток: {total_accumulated_coins:.4f} XAUT. НОВЫЙ ЯКОРЬ: {now_price:.2f} USDT"
                        )
                        print("✅ [ЦИКЛ ЗАКРЫТ] Все тейки исполнены. Обнуляем состояние и открываем новый круг.")

                        initial_reference_price = now_price
                        current_avg_buy_price = now_price
                        current_direction = "BUYING_GRID"
                        total_accumulated_coins = 0.0
                        total_usd_invested = 0.0
                        cycle_usd_spent = 0.0
                        max_drop_reached_in_cycle = 0.0
                        tp_base_volume = 0.0
                        tp_locked_idx = None
                        executed_buy_steps = [False] * len(GRID_BUY_VOLUMES)
                        executed_tp_steps = [False] * 5
                        calculate_reinvest_budget()

                    return  # Одна сделка за тик — даем состоянию обновиться

# ===================================================================================
# ИНИЦИАЛИЗАЦИЯ И ТОЧКА ВХОДА СИСТЕМЫ С WATCHDOG-КОНТРОЛЛЕРОМ
# ===================================================================================
if __name__ == "__main__":
    print("====================================================================")
    print("🤖 ЗАПУСК: АВТОНОМНЫЙ НЕУБИВАЕМЫЙ СЕТОЧНЫЙ РОБОТ ДЛЯ Binance (XAUT-V3)")
    print("====================================================================")
    
    # 1. Запрашиваем параметры округления лимитов лотов XAUT с биржи перед стартом
    FILTERS = get_symbol_filters()
    
    # 2. Синхронизируем состояние с диском (если файл есть — восстановит, если нет — создаст)
    restore_state_from_excel()
    calculate_reinvest_budget()

    # Задаем базовый ориентир цены XAUT для красивой визуализации планов в консоли
    reference_for_display = current_avg_buy_price if current_avg_buy_price > 0 else 42.0 
    
    print(f"\n📊 ТЕКУЩИЙ СТАТУС РОБОТА В ОЗУ: {current_direction}")
    print(f"💰 РАБОЧИЙ БЮДЖЕТ ТЕКУЩЕГО ЦИКЛА: {current_cycle_budget:.2f} USDT")
    print(f"📦 УДЕРЖИВАЕМАЯ ПОЗИЦИЯ: {total_accumulated_coins:.4f} XAUT (Средняя цена: {current_avg_buy_price:.2f} USDT)")
    print(f"📈 Пиковое падение монеты, зафиксированное в текущем круге: {max_drop_reached_in_cycle * 100:.2f}%")
    
    print(f"\n📊 ОРИЕНТИРОВОЧНЫЙ ПЛАН СЕТКИ ПОДКУПОВ (Опора: {reference_for_display:.2f} USDT):")
    print("-" * 95)
    for idx in range(len(GRID_BUY_DROPS)):
        is_bought = executed_buy_steps[idx] if idx < len(executed_buy_steps) else False
        status_lbl = " [УЖЕ ВЫКУПЛЕН]" if is_bought else ""
        print(f"  ⚡ Купить Шаг {idx+1} | При падении до: {reference_for_display * (1 - GRID_BUY_DROPS[idx]):.2f} USDT{status_lbl}")
    print("-" * 95)
    for idx in range(len(GRID_SELL_RISES)):
        is_sold = executed_sell_steps[idx] if idx < len(executed_sell_steps) else False
        status_lbl = " [УЖЕ ПРОДАН]" if is_sold else ""
        print(f"  ⚡ Продать Шаг {idx+1} | При росте до: {reference_for_display * (1 + GRID_SELL_RISES[idx]):.2f} USDT{status_lbl}")
    print("-" * 95)

    ws_client = None

    # Главный внешний цикл, отвечающий за авто-реанимацию при обрывах интернета
    while True:
        try:
            print("\n🔌 Открываем официальный WebSocket поток Binance...")
            # ВАШ ОРИГИНАЛЬНЫЙ РАБОЧИЙ КОНСТРУКТОР СОКЕТА
            ws_client = SpotWebsocketClient(on_message=handle_ticker_message)
            
            # ВАШ ОРИГИНАЛЬНЫЙ МЕТОД ПОДПИСКИ: Используем mini_ticker под монету XAUT!
            # Передаем символ строго в нижнем регистре: symbol="xautusdt"
            ws_client.mini_ticker(symbol=SYMBOL.lower(), id=1)
            
            print(f"✅ Бот успешно зафиксирован в оперативной памяти! Слушаем поток {SYMBOL}")
            last_websocket_packet_time = time.time()

            # Внутренний цикл ежесекундной обработки торговой логики
            while True:
                time.sleep(1) # Проверяем рынок каждую 1 секунду для точности на импульсах
                now_t = time.time()
                
                if current_websocket_price is not None:
                    # 1. Крутим торговую математику в памяти (Ваша каскадная стратегия из Части 7)
                    run_trading_strategy_step(current_websocket_price)
                    
                    # 2. ОЖИВЛЯЕМ КОНСОЛЬ: Минутный отчет на основе текущего режима бота
                    if now_t - last_log_time >= 60:
                        last_log_time = now_t
                        time_str = datetime.now().strftime("%H:%M:%S")
                        
                        if current_direction == "FIRST_RUN_SELL_GRID":
                            ref_p = initial_reference_price if initial_reference_price else current_websocket_price
                            # Берем первый элемент сетки для корректного расчета
                            target_sell = ref_p * (1 + GRID_SELL_RISES[0])
                            deviation = ((current_websocket_price - target_sell) / target_sell) * 100
                            print(f"⏱️ [{time_str}] XAUT: {current_websocket_price:.2f} USDT | РЕЖИМ: SELL-GRID | До цели №1: {deviation:+.2f}%")
                            
                        elif current_direction == "WAIT_MARKET_DROP":
                            ref_p = initial_reference_price if initial_reference_price else current_websocket_price
                            target_drop = ref_p * (1 - MARKET_REBOUND_PERCENT)
                            deviation = ((current_websocket_price - target_drop) / target_drop) * 100
                            print(f"⏱️ [{time_str}] XAUT: {current_websocket_price:.2f} USDT | РЕЖИМ: ЗАСАДА (WAIT_DROP) | До закупа: {deviation:+.2f}%")
                            
                        elif current_direction == "BUYING_GRID":
                            deviation = ((current_websocket_price - current_avg_buy_price) / current_avg_buy_price) * 100 if current_avg_buy_price > 0 else 0.0
                            done_steps = executed_buy_steps.count(True)
                            print(f"⏱️ [{time_str}] XAUT: {current_websocket_price:.2f} USDT | РЕЖИМ: НАКОПЛЕНИЕ (BUY) | Исполнено: {done_steps}/{len(GRID_BUY_VOLUMES)} | От средней: {deviation:+.2f}%")
                            
                        elif current_direction == "WAIT_TAKE_PROFIT":
                            # ✅ ПАТЧ №7: используем ЗАФИКСИРОВАННЫЙ сценарий, если каскад уже идет,
                            # и единую функцию выбора вместо дублирующего блока.
                            target_grid_idx = tp_locked_idx if tp_locked_idx is not None else select_cascade_index(max_drop_reached_in_cycle)

                            parts_count = GRID_PART_PROFIT[target_grid_idx]
                            base_profit_step = GRID_PROFIT[target_grid_idx]
                            
                            next_tp_step = executed_tp_steps.count(True) + 1
                            if next_tp_step > parts_count: 
                                next_tp_step = parts_count
                                
                            target_tp = current_avg_buy_price * (1 + (next_tp_step * base_profit_step))
                            deviation = ((current_websocket_price - target_tp) / target_tp) * 100
                            print(f"⏱️ [{time_str}] XAUT: {current_websocket_price:.2f} USDT | РЕЖИМ: КАСКАД ТЕЙКОВ | Ближайшая цель: №{next_tp_step}/{parts_count} | До фиксации: {deviation:+.2f}%")
     
                # Защитный вочдог от зависания сети (60 секунд отсутствия пакетов от Binance)
                if now_t - last_websocket_packet_time > 60:
                    print("\n🚨 [⚠️ WATCHDOG ТРЕВОГА] Вебсокет Binance завис! Сброс и реконнект...")
                    current_websocket_price = None  
                    try: 
                        ws_client.stop() # Остановка фонового потока
                    except: 
                        pass
                    break # Выходим во внешний цикл для автоматического переподключения

        except KeyboardInterrupt:
            print("\n[СТОП] Работа торгового робота успешно остановлена пользователем.")
            try: 
                ws_client.stop()
            except: 
                pass
            sys.exit(0)
            
        except Exception as socket_error:
            print(f"⚠️ Сетевой сбой сокета Binance: {socket_error}. Повтор подключения через 5 секунд...")
            current_websocket_price = None
            # ✅ ПАТЧ: гасим старый клиент перед реконнектом.
            # Раньше новый сокет создавался поверх старого → утечка потоков и дубли пакетов.
            try:
                if ws_client is not None:
                    ws_client.stop()
            except Exception:
                pass
            ws_client = None
            time.sleep(5)
